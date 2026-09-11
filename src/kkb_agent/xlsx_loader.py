"""M1-10 bounded local XLSX loading; no profile, registry, engine or network API.

The caller supplies an already timeout-controlled binary stream and trusted source
identity. No filename/locator is opened. Only successful tables receive a generated
opaque data reference; storing/resolving it remains the caller's responsibility.
"""

from __future__ import annotations

import io
import json
import math
import multiprocessing
import os
import sys
import tempfile
import time as clock
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, BinaryIO, Literal
from uuid import uuid4

from kkb_agent._xlsx_package import XlsxRejected, inspect_package
from kkb_agent.core import DataContract, SourceRef
from kkb_agent.ingestion_policy import (
    DEFAULT_LIMITS,
    AdmissionError,
    AdmissionLimits,
    FileKind,
    admit_upload,
)
from kkb_agent.xlsx_sheets import select_sheet

type CellValue = str | int | float | bool | datetime | date | time | timedelta | None
#: A cell as the worker writes it to JSON: plain values as they are, temporal
#: values as {"type": ..., "value": ...}.
type EncodedCell = str | int | float | bool | dict[str, Any] | None


@dataclass(frozen=True)
class XlsxTable:
    data_ref: str
    contract: DataContract = field(repr=False)
    selection: str = field(repr=False)
    headers: tuple[CellValue, ...] = field(repr=False)
    rows: tuple[tuple[CellValue, ...], ...] = field(repr=False)

    @property
    def column_ids(self) -> tuple[str, ...]:
        return tuple(f"c{i}" for i in range(len(self.headers)))

    @property
    def row_count(self) -> int:
        return len(self.rows)


@dataclass(frozen=True)
class XlsxLoadResult:
    """Loading outcome only; `loaded` is not an M1-01 `ok` profile."""

    status: Literal["loaded", "needs_input", "rejected"]
    source: SourceRef = field(repr=False)
    code: str | None = None
    sheets: tuple[str, ...] = field(default=(), repr=False)
    table: XlsxTable | None = field(default=None, repr=False)


def _encode_cell(value: object, limits: AdmissionLimits, *, header: bool = False) -> EncodedCell:
    if isinstance(value, str):
        limits.check_field(value, label=header)
    if isinstance(value, float) and not math.isfinite(value):
        raise XlsxRejected("MALFORMED_TABLE")
    if isinstance(value, datetime | date | time):
        return {"type": type(value).__name__, "value": value.isoformat()}
    if isinstance(value, timedelta):
        return {"type": "timedelta", "value": value.total_seconds()}
    # Exact types only; the isinstance only lets the checker see what passed.
    if value is None or (
        isinstance(value, str | int | float | bool) and type(value) in {str, int, float, bool}
    ):
        return value
    raise XlsxRejected("UNSUPPORTED_CELL")


def _decode_cell(value: EncodedCell) -> CellValue:
    if not isinstance(value, dict):
        return value
    constructors: dict[str, type[datetime] | type[date] | type[time]] = {
        "datetime": datetime,
        "date": date,
        "time": time,
    }
    if value["type"] == "timedelta":
        return timedelta(seconds=value["value"])
    return constructors[value["type"]].fromisoformat(value["value"])


def _parse(content: bytes, selection: str | None, limits: AdmissionLimits) -> dict[str, object]:
    # Import in the worker; input XML is validated before the library opens it.
    from openpyxl import load_workbook

    sizes = inspect_package(content, limits)
    with warnings.catch_warnings():
        warnings.simplefilter("error")
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_vba=False, keep_links=False
        )
        try:
            choice = select_sheet(workbook, selection)
            if choice.status != "selected":
                return {"status": choice.status, "code": choice.code, "sheets": choice.sheets}
            selected = choice.selection
            if selected is None:  # select_sheet names the sheet whenever it selects one
                raise XlsxRejected("PARSER_FAILED")
            height, width = sizes[selected]
            if height < 2 or not width:
                raise XlsxRejected("EMPTY_TABLE")
            limits.check_table_size(rows=height - 1, columns=width)
            encoded = []
            output_size = 0
            for ordinal, row in enumerate(
                workbook[selected].iter_rows(min_row=1, max_row=height, min_col=1, max_col=width)
            ):
                limits.check_table_size(rows=ordinal, columns=width)
                if any(cell.data_type in {"f", "e"} for cell in row):
                    raise XlsxRejected("UNSUPPORTED_CELL")
                values = [_encode_cell(cell.value, limits, header=ordinal == 0) for cell in row]
                output_size += len(json.dumps(values, ensure_ascii=True, allow_nan=False).encode())
                if output_size > limits.max_expanded_bytes:
                    raise XlsxRejected("LIMIT_EXCEEDED")
                encoded.append(values)
            if not any(value is not None for value in encoded[0]):
                raise XlsxRejected("MALFORMED_TABLE")
            if not any(value is not None for row in encoded[1:] for value in row):
                raise XlsxRejected("EMPTY_TABLE")
            return {
                "status": "loaded",
                "selection": selected,
                "sheets": choice.sheets,
                "headers": encoded[0],
                "rows": encoded[1:],
            }
        finally:
            workbook.close()


def _write_result(path: Path, result: dict[str, object], limit: int) -> None:
    # Private generated path, never selected by a document. Bound serialization too.
    size = 0
    with open(os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600), "wb") as output:
        for chunk in json.JSONEncoder(ensure_ascii=True, allow_nan=False).iterencode(result):
            data = chunk.encode("utf-8")
            size += len(data)
            if size > limit:
                raise XlsxRejected("LIMIT_EXCEEDED")
            output.write(data)


def _worker(content: bytes, selection: str | None, limits: AdmissionLimits, output: Path) -> None:
    try:
        if sys.platform == "linux":
            import resource

            memory = max(512 * 1024 * 1024, limits.max_expanded_bytes * 8)
            resource.setrlimit(resource.RLIMIT_AS, (memory, memory))
        result = _parse(content, selection, limits)
        _write_result(output, result, limits.max_expanded_bytes)
    except (AdmissionError, XlsxRejected) as error:
        output.unlink(missing_ok=True)
        _write_result(output, {"status": "rejected", "code": str(error.code)}, 1024)
    except MemoryError:
        output.unlink(missing_ok=True)
        _write_result(output, {"status": "rejected", "code": "LIMIT_EXCEEDED"}, 1024)
    except Exception:
        # Invalid ZIP/XML, library errors and warnings must not leak source payloads.
        output.unlink(missing_ok=True)
        _write_result(output, {"status": "rejected", "code": "FORMAT_MISMATCH"}, 1024)


def _run_worker(content: bytes, selection: str | None, limits: AdmissionLimits) -> dict[str, Any]:
    context = multiprocessing.get_context("spawn")
    with tempfile.TemporaryDirectory(prefix="kkb-xlsx-") as directory:
        output = Path(directory) / "result.json"
        process = context.Process(target=_worker, args=(content, selection, limits, output))
        deadline = clock.monotonic() + limits.max_parser_seconds
        try:
            process.start()
            process.join(max(0, deadline - clock.monotonic()))
            if process.is_alive():
                process.terminate()
                process.join(0.2)
                if process.is_alive():
                    process.kill()
                    process.join()
                return {"status": "rejected", "code": "LIMIT_EXCEEDED"}
            if process.exitcode != 0 or not output.is_file():
                return {"status": "rejected", "code": "PARSER_FAILED"}
            with output.open("rb") as stream:
                data = stream.read(max(limits.max_expanded_bytes, 1024) + 1)
            if len(data) > max(limits.max_expanded_bytes, 1024):
                return {"status": "rejected", "code": "LIMIT_EXCEEDED"}
            payload: dict[str, Any] = json.loads(data)  # written by _worker above
            return payload
        finally:
            if process.is_alive():
                process.kill()
                process.join()
            process.close()


def load_xlsx(
    stream: BinaryIO,
    *,
    filename: str,
    content_type: str,
    source: SourceRef,
    selection: str | None = None,
    limits: AdmissionLimits = DEFAULT_LIMITS,
) -> XlsxLoadResult:
    """Admit and load one worksheet; first row is header, blanks stay None.

    Sparse gaps inside actual cell coordinates remain blank rows/cells; declared
    dimensions do not add/drop data. Duplicate/empty labels retain ordinal IDs.
    Text dates/numbers are never coerced. Native Excel date/time styles are read
    by openpyxl; business semantics remain unknown. Formulas/caches, merged cells,
    external relationships and cell errors are rejected rather than approximated.

    The spawn worker is terminated/reaped on timeout; Linux additionally caps its
    address space. Actual expanded bytes and serialized table output each use
    max_expanded_bytes. No workspace persistence or M1-01 full profile is added.
    Applications using spawn must use the usual `if __name__ == '__main__'` guard.
    """
    source = SourceRef.model_validate(source)
    if selection is not None and not isinstance(selection, str):
        raise TypeError("selection must be text or None")
    if selection is not None and len(selection) > limits.max_label_chars:
        return XlsxLoadResult("rejected", source, "LIMIT_EXCEEDED")
    try:
        admitted = admit_upload(stream, filename=filename, content_type=content_type, limits=limits)
        if admitted.kind != FileKind.XLSX:
            return XlsxLoadResult("rejected", source, "UNSUPPORTED_FORMAT")
        result = _run_worker(admitted.content, selection, limits)
    except AdmissionError as error:
        return XlsxLoadResult("rejected", source, error.code.value)
    except OSError:
        return XlsxLoadResult("rejected", source, "PARSER_FAILED")
    if result["status"] != "loaded":
        return XlsxLoadResult(
            result["status"], source, result["code"], tuple(result.get("sheets", ()))
        )
    selected = result["selection"]
    # JSON tuple encoding is unambiguous even when source IDs/titles contain delimiters.
    fragment = SourceRef(
        source_id=json.dumps([source.source_id, selected], ensure_ascii=True),
        source_type=source.source_type,
        locator=json.dumps([source.locator, selected], ensure_ascii=True),
        retrieved_at=source.retrieved_at,
        content_hash=source.content_hash,
    )
    data_ref = f"data:{uuid4().hex}"
    table = XlsxTable(
        data_ref,
        DataContract(contract_id=data_ref, source=fragment, title=selected),
        selected,
        tuple(_decode_cell(value) for value in result["headers"]),
        tuple(tuple(_decode_cell(value) for value in row) for row in result["rows"]),
    )
    return XlsxLoadResult("loaded", source, sheets=tuple(result["sheets"]), table=table)
