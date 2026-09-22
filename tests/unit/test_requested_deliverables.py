"""Explicit deliverables survive execution, publication, API projection and restart."""

import io
import json

import pytest
from openpyxl import Workbook
from test_conversation_persistence import launch as launch  # noqa: F401
from test_conversation_persistence import open_conversation, owner
from test_document_table import ask

from kkb_agent.agent import Workspace
from kkb_agent.lakehouse import open_lakehouse

RAW = [0.476, 0.49, 0.5, 0.505, 0.51, 0.515, *([0.52] * 6)]


def workbook(*, mixed=False, header="Faiz"):
    book = Workbook()
    sheet = book.active
    sheet.append(["Ay", header])
    for month, number in enumerate(RAW, 1):
        sheet.append([f"2024-{month:02d}", number])
        sheet.cell(month + 1, 2).number_format = "0.0%"
    if mixed:
        sheet.cell(3, 2).number_format = "General"
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def plan(operation="mean", *, present=True):
    read = {
        "step_id": "read",
        "tool": "web_url",
        "present": present,
        "arguments": {"attached": True, "period_column": "Ay", "value_columns": ["Faiz"]},
    }
    arguments = {
        "operation": "summarize",
        "source": "read",
        "column": "Faiz",
        "aggregation": operation,
        "requested_aggregation": operation,
    }
    if operation == "change":
        arguments = {
            "operation": "change",
            "source": "read",
            "column": "Faiz",
            "kind": "difference",
        }
    return json.dumps(
        {
            "question": "Belgeden tablo ve analiz",
            "interpret": False,
            "steps": [
                read,
                {
                    "step_id": "derive",
                    "tool": "table_transform",
                    "arguments": arguments,
                    "depends_on": ["read"],
                },
            ],
        }
    )


def uploaded(client, sid, headers, content=None):
    response = client.post(
        f"/v1/sessions/{sid}/documents",
        content=workbook() if content is None else content,
        headers={
            **headers,
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()["document_ref"]


@pytest.mark.parametrize(
    "operation,question",
    [
        (
            "mean",
            "Bu dosyadaki aylık faiz oranlarını tablo halinde göster ve 2024 ortalamasını hesapla.",
        ),
        ("max", "Tabloyu göster ve en yüksek değeri söyle."),
        ("change", "Tabloyu göster ve değişimi hesapla."),
    ],
)
def test_requested_table_and_analysis_survive_restart_and_selection(launch, operation, question):
    runner, client = launch(plan(operation), "Tablo ve hesaplama hazır.")
    token = owner(client)
    sid, headers, _ = open_conversation(client, token)
    document = uploaded(client, sid, headers)
    turn = ask(client, sid, headers, question, document_ref=document)
    assert turn["answer"]["status"] in ("accepted", "qualified"), turn
    published = turn["answer"]["result"]
    assert published is not None, turn
    if operation in ("mean", "max"):
        writer_context = runner._provider.seen[-1][1].content
        assert "published in an earlier turn" not in writer_context
        assert writer_context.count("this turn's result") == 2
    original = [cell for cell in published["cells"] if cell["label"] == "Faiz"]
    assert len(original) == 12
    assert [cell["value"] for cell in original] == pytest.approx([x * 100 for x in RAW])
    assert all(cell["measure"] == "rate" and cell["unit"] == "percent" for cell in original)
    related = published["related"]
    if operation == "change":
        # The transformed table already contains the original column: no duplicate full table.
        assert not related
        assert len(published["cells"]) == 23
    else:
        assert len(related) == 1
        assert related[0]["cells"][0]["value"] == pytest.approx(
            sum(RAW) * 100 / 12 if operation == "mean" else max(RAW) * 100
        )
    for artifact in [published, *related]:
        assert artifact["proof"]["availability"] == "available"
        assert artifact["proof"]["steps"][0]["result_version"] == artifact["version"]
    _, reopened = launch()
    restored = reopened.get(f"/v1/conversations/{sid}/results/{turn['turn_id']}", headers=headers)
    assert restored.status_code == 200, restored.text
    assert restored.json() == published
    # Default follow-up selects the preserved monthly table, never the one-row summary.
    from kkb_agent.api.service import SelectedResult

    selected = runner.resolve_selection(
        sid, SelectedResult(turn["turn_id"], published["result_ref"], published["version"])
    )
    assert len([v for v in selected.value.values if v.label == "Faiz"]) == 12


def test_invalid_rate_sum_is_terminal_even_with_runtime_repair_enabled(launch):
    runner, client = launch(plan("sum"), plan("mean"), "Ortalama hazır.")
    sid, headers, _ = open_conversation(client, owner(client))
    document = uploaded(client, sid, headers)
    turn = ask(
        client, sid, headers, "2024 faiz oranlarının toplamını hesapla.", document_ref=document
    )
    assert turn["answer"]["status"] == "abstained", turn
    assert turn["answer"]["needs_clarification"] is False
    assert turn["answer"]["result"] is None
    assert "Başka bir metrik hesaplanmadı" in turn["answer"]["text"]
    assert len(runner._provider.seen) == 1
    with open_lakehouse(runner.run_root / sid / "lake") as lake:
        assert Workspace(lake.connection, sid).names() == []
        assert lake.connection.execute("SELECT count(*) FROM gold_cells").fetchone()[0] == 0
    requested = ask(
        client, sid, headers, "O zaman 2024 ortalamasını hesapla.", document_ref=document
    )
    assert requested["answer"]["result"] is not None, requested
    assert requested["answer"]["result"]["related"][0]["cells"][0]["value"] == pytest.approx(
        50.9666666667
    )


def test_summary_alone_does_not_publish_unrequested_upstream(launch):
    _, client = launch(plan("mean", present=False), "Ortalama hazır.")
    sid, headers, _ = open_conversation(client, owner(client))
    document = uploaded(client, sid, headers)
    turn = ask(client, sid, headers, "2024 ortalamasını hesapla.", document_ref=document)
    assert len(turn["answer"]["result"]["cells"]) == 1, turn
    assert not turn["answer"]["result"]["related"]


def test_planner_cannot_replace_requested_sum_with_mean_in_its_initial_plan(launch, monkeypatch):
    proposed = json.loads(plan("mean"))
    proposed["steps"][-1]["arguments"]["requested_aggregation"] = "sum"
    proposed["steps"][-1]["assumption"] = "Oran toplanamadığı için ortalama hesaplandı."

    def forbidden_calculation(*args, **kwargs):
        raise AssertionError("Substituted aggregation must never execute")

    monkeypatch.setattr("kkb_agent.tools.table_transform.summarize_column", forbidden_calculation)
    runner, client = launch(json.dumps(proposed), plan("mean"))
    sid, headers, _ = open_conversation(client, owner(client))
    turn = ask(
        client,
        sid,
        headers,
        "2024 faiz oranlarının toplamını hesapla.",
        document_ref=uploaded(client, sid, headers),
    )
    assert turn["answer"]["status"] == "abstained", turn
    assert turn["answer"]["result"] is None
    assert "Alternatif hesaplama yapılmadı" in turn["answer"]["text"]
    assert len(runner._provider.seen) == 1


def test_requested_metric_cannot_be_defaulted_to_the_proposed_metric():
    from pydantic import ValidationError

    from kkb_agent.tools.table_transform import TableTransformArguments

    arguments = json.loads(plan("mean"))["steps"][-1]["arguments"]
    del arguments["requested_aggregation"]
    with pytest.raises(ValidationError, match="requested_aggregation"):
        TableTransformArguments.model_validate(arguments)


def test_equal_cells_from_independent_documents_keep_both_proofs(launch):
    first = json.loads(plan())["steps"][0]
    first["arguments"]["document"] = "belge_1"
    second = {
        **first,
        "step_id": "other",
        "arguments": {
            **first["arguments"],
            "document": "belge_2",
        },
    }
    _, client = launch(
        json.dumps(
            {
                "question": "İki belgenin tablolarını göster",
                "interpret": False,
                "steps": [first, second],
            }
        ),
        "Her iki tablo hazır.",
    )
    sid, headers, _ = open_conversation(client, owner(client))
    uploaded(client, sid, headers)
    document = uploaded(client, sid, headers, workbook(header="Faiz (%)"))
    turn = ask(client, sid, headers, "İki belgenin tablolarını göster.", document_ref=document)
    result = turn["answer"]["result"]
    assert result is not None, turn
    assert len(result["related"]) == 1
    other = result["related"][0]
    assert [cell["value"] for cell in result["cells"]] == [cell["value"] for cell in other["cells"]]
    assert result["proof"]["availability"] == other["proof"]["availability"] == "available"
    assert result["version"] != other["version"]
