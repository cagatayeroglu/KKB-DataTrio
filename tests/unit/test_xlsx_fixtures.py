"""Exact synthetic fixture oracles and M1-09 sheet-selection behavior."""

import subprocess
import sys
from contextlib import closing
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from kkb_agent.ingestion_policy import FileKind, admit_upload
from kkb_agent.xlsx_sheets import select_sheet

ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / "tests/fixtures/ingestion"
OBSERVATIONS = [
    ("as_of", "measure", "label"),
    (datetime(2026, 1, 31), 100, "Örnek"),
    (datetime(2026, 2, 28), None, "çığ"),
    (datetime(2026, 3, 31), 125.5, "C"),
]
CATEGORIES = [("code", "description"), ("A", "Synthetic alpha"), ("B", "Örnek beta")]


def open_fixture(filename):
    return load_workbook(
        FIXTURES / filename, read_only=True, data_only=False, keep_vba=False, keep_links=False
    )


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("single_sheet.xlsx", {"Observations": OBSERVATIONS}),
        ("multi_sheet.xlsx", {"Observations": OBSERVATIONS, "Categories": CATEGORIES}),
    ],
)
def test_synthetic_xlsx_content_and_package(filename, expected):
    path = FIXTURES / filename
    assert 0 < path.stat().st_size < 10_000
    content = path.read_bytes()
    admitted = admit_upload(
        BytesIO(content),
        filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert admitted.kind is FileKind.XLSX
    assert admitted.content == content
    with ZipFile(BytesIO(content)) as archive:
        assert archive.testzip() is None
        types = archive.read("[Content_Types].xml")
        assert (
            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml" in types
        )
        assert b"macroEnabled" not in types
        assert not any(
            name.lower().endswith("vbaproject.bin")
            or name.lower().startswith(("xl/externallinks/", "xl/embeddings/"))
            for name in archive.namelist()
        )
        assert all(entry.date_time == (2026, 1, 1, 0, 0, 0) for entry in archive.infolist())
    with closing(open_fixture(filename)) as workbook:
        assert workbook.sheetnames == list(expected)
        assert workbook.template is False
        assert workbook.properties.created == datetime(2026, 1, 1)
        assert workbook.properties.modified == datetime(2026, 1, 1)
        for sheet in workbook.worksheets:
            assert sheet.sheet_state == "visible"
            assert (sheet.max_row, sheet.max_column) == (
                len(expected[sheet.title]),
                len(expected[sheet.title][0]),
            )
            assert list(sheet.values) == expected[sheet.title]
            assert all(cell.data_type != "f" for row in sheet for cell in row)
        assert workbook["Observations"]["A2"].number_format == "yyyy-mm-dd"
        assert workbook["Observations"]["B3"].value is None


@pytest.mark.parametrize("selection", [None, "Observations"])
def test_single_sheet_can_be_selected_unambiguously(selection):
    with closing(open_fixture("single_sheet.xlsx")) as workbook:
        result = select_sheet(workbook, selection)
        assert result.status == "selected"
        assert result.sheets == ("Observations",)
        assert result.selection == "Observations"
        assert result.code is None


def test_multiple_sheets_need_selection_even_when_active_sheet_changes():
    with closing(open_fixture("multi_sheet.xlsx")) as workbook:
        workbook.active = 1
        result = select_sheet(workbook)
        assert result.status == "needs_input"
        assert result.sheets == ("Observations", "Categories")
        assert result.selection is None
        assert result.code == "SELECTION_REQUIRED"


@pytest.mark.parametrize("selection", ["Observations", "Categories"])
def test_explicit_multi_sheet_selection(selection):
    with closing(open_fixture("multi_sheet.xlsx")) as workbook:
        result = select_sheet(workbook, selection)
        assert result.status == "selected"
        assert result.sheets == ("Observations", "Categories")
        assert result.selection == selection
        assert result.code is None


@pytest.mark.parametrize("filename", ["single_sheet.xlsx", "multi_sheet.xlsx"])
@pytest.mark.parametrize("selection", ["Missing", "observations", " Observations ", ""])
def test_invalid_selection_is_structured_without_fallback(filename, selection):
    with closing(open_fixture(filename)) as workbook:
        result = select_sheet(workbook, selection)
        assert result.status == "needs_input"
        assert result.selection is None
        assert result.code == "UNKNOWN_SHEET"
        assert result.sheets == tuple(workbook.sheetnames)


def test_hidden_worksheet_does_not_silently_remove_ambiguity():
    with closing(Workbook()) as workbook:
        workbook.active.title = "Visible"
        workbook.create_sheet("Hidden").sheet_state = "hidden"
        assert select_sheet(workbook).code == "SELECTION_REQUIRED"
        assert select_sheet(workbook).sheets == ("Visible", "Hidden")
        assert select_sheet(workbook, "Hidden").selection == "Hidden"


def test_chart_tabs_are_not_tabular_worksheets():
    with closing(Workbook()) as workbook:
        workbook.remove(workbook.active)
        workbook.create_chartsheet("Chart")
        result = select_sheet(workbook)
        assert result.status == "rejected"
        assert result.code == "NO_WORKSHEETS"
        assert result.selection is None
        assert result.sheets == ()


def test_selection_does_not_iterate_cells(monkeypatch):
    def unexpected_read(*args, **kwargs):
        pytest.fail("Sheet selection must not parse cells")

    with closing(open_fixture("multi_sheet.xlsx")) as workbook:
        for sheet in workbook.worksheets:
            monkeypatch.setattr(sheet, "iter_rows", unexpected_read)
        assert select_sheet(workbook).code == "SELECTION_REQUIRED"
        assert select_sheet(workbook, "Categories").selection == "Categories"


def test_generator_is_deterministic(tmp_path):
    command = [sys.executable, str(ROOT / "scripts/generate_xlsx_fixtures.py"), str(tmp_path)]
    subprocess.run(command, check=True, capture_output=True)
    first = {path.name: path.read_bytes() for path in tmp_path.glob("*.xlsx")}
    assert set(first) == {"single_sheet.xlsx", "multi_sheet.xlsx"}
    subprocess.run(command, check=True, capture_output=True)
    assert first == {path.name: path.read_bytes() for path in tmp_path.glob("*.xlsx")}
    for filename in first:
        with closing(load_workbook(tmp_path / filename, read_only=True)) as workbook:
            assert list(workbook["Observations"].values) == OBSERVATIONS
