"""Small synthetic parser/admission/resource regressions for M1-10."""

import io
import json
import multiprocessing
import os
import time
from dataclasses import replace
from datetime import UTC, date, datetime, timedelta
from datetime import time as time_of_day
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import pytest
from openpyxl import Workbook

from kkb_agent import _xlsx_package as package
from kkb_agent import xlsx_loader as loader
from kkb_agent.core import Frequency, SourceRef, StockFlow
from kkb_agent.ingestion_policy import DEFAULT_LIMITS, MIME_TYPES, AdmissionError, FileKind

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures/ingestion"
SINGLE = (FIXTURES / "single_sheet.xlsx").read_bytes()
MULTI = (FIXTURES / "multi_sheet.xlsx").read_bytes()
SHEET = "xl/worksheets/sheet1.xml"
SOURCE = SourceRef(
    source_id="source:synthetic",
    source_type="upload",
    locator="opaque:fixture",
    retrieved_at=datetime(2026, 9, 8, tzinfo=UTC),
    content_hash="synthetic-hash",
)


def workbook_bytes(rows, *, second_sheet=False, text_cells=()):
    workbook = Workbook()
    workbook.active.title = "Data"
    for row in rows:
        workbook.active.append(row)
    for coordinate in text_cells:
        workbook.active[coordinate].data_type = "s"
    if second_sheet:
        workbook.create_sheet("Other").append(["label"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def rewrite(content, updates, *, compression=ZIP_DEFLATED):
    output = io.BytesIO()
    with ZipFile(io.BytesIO(content)) as original, ZipFile(output, "w", compression) as result:
        for name in original.namelist():
            data = original.read(name)
            if name in updates:
                update = updates[name]
                data = update(data) if callable(update) else update
            if data is not None:
                result.writestr(name, data)
        for name, data in updates.items():
            if name not in original.namelist():
                result.writestr(name, data)
    return output.getvalue()


def load(content=SINGLE, **kwargs):
    return loader.load_xlsx(
        io.BytesIO(content),
        filename="synthetic.xlsx",
        content_type=MIME_TYPES[FileKind.XLSX],
        source=SOURCE,
        **kwargs,
    )


@pytest.mark.parametrize("selection", [None, "Observations"])
def test_public_loader_exact_fixture_unknowns_provenance_and_no_payload_repr(selection):
    result = load(selection=selection)
    assert result.status == "loaded" and result.code is None
    table = result.table
    assert table.headers == ("as_of", "measure", "label")
    assert table.column_ids == ("c0", "c1", "c2")
    assert table.row_count == 3
    assert table.rows == (
        (datetime(2026, 1, 31), 100, "Örnek"),
        (datetime(2026, 2, 28), None, "çığ"),
        (datetime(2026, 3, 31), 125.5, "C"),
    )
    assert result.source == SOURCE
    contract = table.contract
    assert contract.contract_id == table.data_ref and table.data_ref.startswith("data:")
    assert json.loads(contract.source.source_id) == [SOURCE.source_id, "Observations"]
    assert json.loads(contract.source.locator) == [SOURCE.locator, "Observations"]
    assert contract.source.retrieved_at == SOURCE.retrieved_at
    assert contract.source.content_hash == SOURCE.content_hash
    assert contract.frequency is Frequency.UNKNOWN and contract.stock_flow is StockFlow.UNKNOWN
    assert contract.unit is contract.cumulative is contract.reset_behavior is None
    assert contract.time_column is contract.time_semantics is None
    assert not contract.transformations
    assert "Örnek" not in repr(result) + repr(table)
    assert SOURCE.locator not in repr(result) + repr(table)


@pytest.mark.parametrize(
    "selection,code", [(None, "SELECTION_REQUIRED"), ("missing", "UNKNOWN_SHEET")]
)
def test_public_ambiguous_selection_never_publishes_a_table(selection, code):
    result = load(MULTI, selection=selection)
    assert result.status == "needs_input" and result.code == code
    assert result.table is None
    assert result.sheets == ("Observations", "Categories")


def test_selected_sheets_have_distinct_data_and_source_fragment_identities():
    first, second = load(MULTI, selection="Observations"), load(MULTI, selection="Categories")
    assert first.table.data_ref != second.table.data_ref
    assert first.table.contract.source.source_id != second.table.contract.source.source_id
    assert second.table.headers == ("code", "description")
    assert second.table.rows == (("A", "Synthetic alpha"), ("B", "Örnek beta"))


def test_unknown_text_missing_duplicate_headers_boolean_and_time_are_preserved():
    content = workbook_bytes(
        [
            ["duplicate", "duplicate", None, "cumulative stock USD", "time"],
            ["01/02/2026", "1,234", None, False, timedelta(hours=26)],
            ["NA", "  ", "=instructions as text", True, None],
        ],
        # Explicit text stays data, regardless of the optional XML writer backend.
        text_cells=("C3",),
    )
    result = load(content)
    assert result.status == "loaded"
    assert result.table.column_ids == ("c0", "c1", "c2", "c3", "c4")
    assert result.table.headers[:3] == ("duplicate", "duplicate", None)
    assert result.table.rows == (
        ("01/02/2026", "1,234", None, False, timedelta(hours=26)),
        ("NA", "  ", "=instructions as text", True, None),
    )
    assert result.table.contract.cumulative is None


@pytest.mark.parametrize("dimension", [b"A1:A1", b"A1:XFD1048576"])
def test_false_declared_dimensions_do_not_truncate_or_expand_table(dimension):
    content = rewrite(SINGLE, {SHEET: lambda b: b.replace(b"A1:C4", dimension)})
    result = loader._parse(content, None, DEFAULT_LIMITS)
    assert len(result["rows"]) == 3 and len(result["headers"]) == 3


def test_sparse_internal_blank_rows_are_preserved_from_actual_coordinates():
    content = workbook_bytes([["header"], ["first"], [None], ["last"]])
    result = loader._parse(content, None, DEFAULT_LIMITS)
    assert result["rows"] == [["first"], [None], ["last"]]


@pytest.mark.parametrize("rows", [[], [["header"]], [["header"], [None]]])
def test_empty_or_header_only_tables_are_rejected(rows):
    with pytest.raises(package.XlsxRejected, match="EMPTY_TABLE"):
        loader._parse(workbook_bytes(rows), None, DEFAULT_LIMITS)


@pytest.mark.parametrize(
    "kwargs",
    [
        {"filename": "../secret.xlsx"},
        {"content_type": "text/csv"},
        {"filename": "synthetic.xlsm"},
        {"limits": replace(DEFAULT_LIMITS, max_bytes=len(SINGLE) - 1)},
        {"limits": replace(DEFAULT_LIMITS, max_archive_entries=1)},
    ],
)
def test_admission_rejects_before_worker_start(monkeypatch, kwargs):
    def forbidden(*args, **kwargs):
        pytest.fail("Rejected input reached worker")

    monkeypatch.setattr(loader, "_run_worker", forbidden)
    args = {
        "filename": "synthetic.xlsx",
        "content_type": MIME_TYPES[FileKind.XLSX],
        "source": SOURCE,
    }
    args.update(kwargs)
    result = loader.load_xlsx(io.BytesIO(SINGLE), **args)
    assert result.status == "rejected" and result.table is None


@pytest.mark.parametrize(
    "changes",
    [
        {"max_rows": 2},
        {"max_columns": 2},
        {"max_cells": 8},
        {"max_field_chars": 4},
        {"max_label_chars": 4},
        {"max_sheets": 1},
    ],
)
def test_parser_dimension_and_field_limits(changes):
    with pytest.raises(AdmissionError, match="LIMIT_EXCEEDED"):
        loader._parse(MULTI, "Observations", replace(DEFAULT_LIMITS, **changes))


def test_exact_table_limits_pass():
    limits = replace(DEFAULT_LIMITS, max_rows=3, max_columns=3, max_cells=9)
    assert len(loader._parse(SINGLE, None, limits)["rows"]) == 3


@pytest.mark.parametrize("replacement", [b'r="A100001"', b'r="XFD2"', b'r="A0"'])
def test_hostile_cell_coordinates_fail_before_library(monkeypatch, replacement):
    import openpyxl

    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: pytest.fail("library reached"))
    content = rewrite(SINGLE, {SHEET: lambda b: b.replace(b'r="A2"', replacement)})
    with pytest.raises((package.XlsxRejected, AdmissionError)):
        loader._parse(content, None, DEFAULT_LIMITS)


@pytest.mark.parametrize(
    "xml",
    [
        b'<!DOCTYPE worksheet [<!ENTITY x "payload">]><worksheet>&x;</worksheet>',
        b'<!DOCTYPE worksheet SYSTEM "file:///not-readable"><worksheet/>',
    ],
)
def test_dtd_and_entities_rejected_before_library(monkeypatch, xml):
    import openpyxl

    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: pytest.fail("library reached"))
    with pytest.raises(package.XlsxRejected, match="UNSAFE_XLSX"):
        loader._parse(rewrite(SINGLE, {SHEET: xml}), None, DEFAULT_LIMITS)


def test_formula_is_rejected_even_with_cached_value_or_on_unselected_sheet():
    def cache_formula(xml):
        for empty in (b"<v />", b"<v/>", b"<v></v>"):
            xml = xml.replace(empty, b"<v>2</v>")
        return xml

    formula = workbook_bytes([["value"], ["=1+1"]], second_sheet=True)
    formula = rewrite(formula, {SHEET: cache_formula})
    with ZipFile(io.BytesIO(formula)) as archive:
        assert b"<f>1+1</f><v>2</v>" in archive.read(SHEET)
    with pytest.raises(package.XlsxRejected, match="FORMULA_UNSUPPORTED"):
        loader._parse(formula, "Other", DEFAULT_LIMITS)


@pytest.mark.parametrize(
    "target,mode",
    [
        ("https://example.invalid/private", "External"),
        ("file:///private", "External"),
        ("../../../outside.xml", "Internal"),
        ("/xl/missing.xml", "Internal"),
        ("https://example.invalid/private", "Internal"),
    ],
)
def test_external_and_invalid_relationships_rejected(target, mode):
    rel = f'<Relationship Id="bad" Type="x" Target="{target}" TargetMode="{mode}"/>'
    content = rewrite(
        SINGLE,
        {
            "xl/_rels/workbook.xml.rels": lambda b: b.replace(
                b"</Relationships>", rel.encode() + b"</Relationships>"
            )
        },
    )
    with pytest.raises(package.XlsxRejected):
        loader._parse(content, None, DEFAULT_LIMITS)


def test_macro_content_type_is_rejected_even_without_vba_member():
    content = rewrite(
        SINGLE,
        {
            "[Content_Types].xml": lambda b: b.replace(
                b"spreadsheetml.sheet.main+xml", b"spreadsheetml.sheet.macroEnabled.main+xml"
            )
        },
    )
    with pytest.raises(package.XlsxRejected, match="UNSAFE_XLSX"):
        loader._parse(content, None, DEFAULT_LIMITS)


def test_duplicate_relationship_ids_rejected():
    content = rewrite(
        SINGLE, {"xl/_rels/workbook.xml.rels": lambda b: b.replace(b'Id="rId2"', b'Id="rId1"')}
    )
    with pytest.raises(package.XlsxRejected, match="FORMAT_MISMATCH"):
        loader._parse(content, None, DEFAULT_LIMITS)


def test_merged_cells_and_excel_errors_are_explicitly_unsupported():
    merged = rewrite(
        SINGLE,
        {
            SHEET: lambda b: b.replace(
                b"</worksheet>", b'<mergeCells><mergeCell ref="A2:B2"/></mergeCells></worksheet>'
            )
        },
    )
    with pytest.raises(package.XlsxRejected, match="MERGED_CELLS_UNSUPPORTED"):
        loader._parse(merged, None, DEFAULT_LIMITS)
    with pytest.raises(package.XlsxRejected, match="UNSUPPORTED_CELL"):
        loader._parse(workbook_bytes([["header"], ["#DIV/0!"]]), None, DEFAULT_LIMITS)


def test_nonfinite_numeric_values_are_rejected():
    content = rewrite(SINGLE, {SHEET: lambda b: b.replace(b"<v>100</v>", b"<v>1e999</v>")})
    with pytest.raises(package.XlsxRejected, match="MALFORMED_TABLE"):
        loader._parse(content, None, DEFAULT_LIMITS)


def test_actual_expansion_budget_stops_stream_before_library(monkeypatch):
    real = package.ZipFile
    reads = []

    class Archive:
        def __enter__(self):
            self.inner = real(io.BytesIO(SINGLE))
            return self

        def __exit__(self, *args):
            self.inner.close()

        def namelist(self):
            return self.inner.namelist()

        def infolist(self):
            return [SimpleNamespace(filename="docProps/app.xml", file_size=1, compress_size=100)]

        def open(self, member):
            stream = self.inner.open(member.filename)
            original = stream.read

            def read(size):
                reads.append(size)
                return original(size)

            stream.read = read
            return stream

    monkeypatch.setattr(package, "ZipFile", lambda *a: Archive())
    with pytest.raises(AdmissionError, match="LIMIT_EXCEEDED"):
        package.inspect_package(SINGLE, replace(DEFAULT_LIMITS, max_expanded_bytes=10))
    assert reads == [11]


@pytest.mark.parametrize("kind", ["corrupt_xml", "bad_crc", "formula"])
def test_public_parser_failures_are_sanitized_and_publish_no_partial_table(kind, capfd):
    if kind == "corrupt_xml":
        content = rewrite(SINGLE, {SHEET: b"<private-payload>broken"})
    elif kind == "bad_crc":
        content = rewrite(SINGLE, {}, compression=ZIP_STORED)
        content = content.replace(b"Observations", b"ObservationX", 1)
    else:
        content = workbook_bytes([["private-payload"], ["=1+1"]])
    result = load(content)
    assert result.status == "rejected" and result.table is None
    assert result.code in {"FORMAT_MISMATCH", "FORMULA_UNSUPPORTED"}
    assert "private-payload" not in repr(result)
    captured = capfd.readouterr()
    assert not captured.out and not captured.err


def _slow_worker(content, selection, limits, output):
    time.sleep(30)


def test_timeout_terminates_and_reaps_parser_and_cleans_private_output(monkeypatch, tmp_path):
    context = multiprocessing.get_context("spawn")
    observed = {}

    def create(**kwargs):
        kwargs["target"] = _slow_worker
        process = context.Process(**kwargs)

        class TrackedProcess:
            def __getattr__(self, name):
                return getattr(process, name)

            def close(self):
                observed["alive_before_close"] = process.is_alive()
                observed["exitcode"] = process.exitcode
                process.close()
                observed["closed"] = True

        return TrackedProcess()

    monkeypatch.setattr(
        loader.multiprocessing, "get_context", lambda *a: SimpleNamespace(Process=create)
    )
    original_temp = loader.tempfile.TemporaryDirectory
    monkeypatch.setattr(
        loader.tempfile, "TemporaryDirectory", lambda **kw: original_temp(dir=tmp_path, **kw)
    )
    start = time.monotonic()
    result = load(limits=replace(DEFAULT_LIMITS, max_parser_seconds=1))
    assert time.monotonic() - start < 4
    assert result.status == "rejected" and result.code == "LIMIT_EXCEEDED"
    assert result.table is None and not list(tmp_path.iterdir())
    assert observed["closed"] and not observed["alive_before_close"]
    assert observed["exitcode"] is not None


def test_serialized_result_cap_and_private_file_mode(tmp_path):
    path = tmp_path / "result.json"
    loader._write_result(path, {"value": 12}, 13)
    if os.name == "posix":  # Windows uses ACLs, not POSIX permission bits.
        assert path.stat().st_mode & 0o777 == 0o600
    assert json.loads(path.read_text()) == {"value": 12}
    with pytest.raises(package.XlsxRejected, match="LIMIT_EXCEEDED"):
        loader._write_result(tmp_path / "too-large.json", {"value": "x" * 100}, 13)


@pytest.mark.parametrize("suffix,encoding", [("xml", "utf-16"), ("bin", "utf-8")])
def test_entities_cannot_bypass_checks_by_encoding_or_part_suffix(suffix, encoding):
    text = f'<?xml version="1.0" encoding="{encoding}"?><!DOCTYPE x [<!ENTITY e "x">]><x>&e;</x>'
    content = rewrite(SINGLE, {f"custom/extra.{suffix}": text.encode(encoding)})
    with pytest.raises(package.XlsxRejected, match="UNSAFE_XLSX"):
        loader._parse(content, None, DEFAULT_LIMITS)


def test_xml_depth_is_bounded():
    content = rewrite(SINGLE, {"custom/extra.xml": b"<x>" * 65 + b"</x>" * 65})
    with pytest.raises(AdmissionError, match="LIMIT_EXCEEDED"):
        loader._parse(content, None, DEFAULT_LIMITS)


def test_shared_string_repetition_cannot_bypass_materialized_output_cap():
    rows = b"".join(
        f'<row r="{i}"><c r="A{i}" t="s"><v>0</v></c></row>'.encode() for i in range(2, 42)
    )
    worksheet = (
        f'<worksheet xmlns="{package.SHEET_NS}"><sheetData>'.encode()
        + b'<row r="1"><c r="A1" t="inlineStr"><is><t>header</t></is></c></row>'
        + rows
        + b"</sheetData></worksheet>"
    )
    strings = f'<sst xmlns="{package.SHEET_NS}"><si><t>{"x" * 1000}</t></si></sst>'.encode()
    override = (
        b'<Override PartName="/xl/sharedStrings.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
    )
    content = rewrite(
        SINGLE,
        {
            SHEET: worksheet,
            "xl/sharedStrings.xml": strings,
            "[Content_Types].xml": lambda b: b.replace(b"</Types>", override + b"</Types>"),
        },
    )
    with ZipFile(io.BytesIO(content)) as archive:
        expanded = sum(info.file_size for info in archive.infolist())
    with pytest.raises(package.XlsxRejected, match="LIMIT_EXCEEDED"):
        loader._parse(content, None, replace(DEFAULT_LIMITS, max_expanded_bytes=expanded))


def _crashed_worker(content, selection, limits, output):
    os._exit(9)


def test_worker_crash_returns_sanitized_failure_without_table(monkeypatch):
    context = multiprocessing.get_context("spawn")

    def create(**kwargs):
        kwargs["target"] = _crashed_worker
        return context.Process(**kwargs)

    monkeypatch.setattr(
        loader.multiprocessing, "get_context", lambda *a: SimpleNamespace(Process=create)
    )
    result = load()
    assert result.status == "rejected" and result.code == "PARSER_FAILED"
    assert result.table is None


def test_successful_worker_also_cleans_private_output(monkeypatch, tmp_path):
    original_temp = loader.tempfile.TemporaryDirectory
    monkeypatch.setattr(
        loader.tempfile, "TemporaryDirectory", lambda **kw: original_temp(dir=tmp_path, **kw)
    )
    assert load().status == "loaded"
    assert not list(tmp_path.iterdir())


def test_blank_sheet_title_fails_safely_before_shared_contract_construction():
    content = rewrite(
        SINGLE, {"xl/workbook.xml": lambda b: b.replace(b'name="Observations"', b'name="   "')}
    )
    result = load(content)
    assert result.status == "rejected" and result.code == "FORMAT_MISMATCH"
    assert result.table is None


class _Text(str):
    """A str subclass, as a library-specific rich-text type can be."""


@pytest.mark.parametrize("plain", ["label", 7, 1.5, True, None])
def test_plain_cell_values_pass_through_unchanged(plain):
    assert loader._encode_cell(plain, DEFAULT_LIMITS) == plain


@pytest.mark.parametrize("unusual", [Decimal("1.5"), _Text("rich"), [1], b"raw"])
def test_only_exact_plain_types_are_accepted(unusual):
    # M0-08 typed this path; the exact-type rule it enforces must not loosen.
    with pytest.raises(package.XlsxRejected, match="UNSUPPORTED_CELL"):
        loader._encode_cell(unusual, DEFAULT_LIMITS)


@pytest.mark.parametrize(
    "temporal",
    [datetime(2026, 1, 31, 12, 30), date(2026, 1, 31), time_of_day(9, 15), timedelta(days=2)],
)
def test_temporal_cells_survive_the_json_round_trip(temporal):
    encoded = loader._encode_cell(temporal, DEFAULT_LIMITS)

    assert loader._decode_cell(json.loads(json.dumps(encoded))) == temporal


def percentage_workbook(number, fmt):
    book = Workbook()
    book.active.append(["Ay", "Faiz (%)"])
    book.active.append(["2024-01", number])
    book.active.append(["2024-02", 48.0])
    book.active["B2"].number_format = fmt
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


@pytest.mark.parametrize(
    "number,fmt,expected",
    [(0.476, "0.0%", 47.6), (47.6, "General", 47.6), ("47.6", "0.0%", "47.6")],
)
def test_percentage_display_is_normalized_once_with_proof(number, fmt, expected):
    result = load(percentage_workbook(number, fmt))
    assert result.status == "loaded"
    assert result.table.rows[0][1] == expected
    from kkb_agent.tools._document_column import column_snapshot, read_column
    from kkb_agent.tools._document_value import grids

    column = read_column(grids(result.table, None), "Ay", "Faiz")
    snapshot = json.loads(column_snapshot(column, SOURCE, MIME_TYPES[FileKind.XLSX]))
    if number == 0.476:
        assert snapshot["normalizations"][0]["operation"] == "excel_percent_points"
    else:
        assert not result.table.normalizations


def test_custom_percentage_format_requires_explicit_support():
    with pytest.raises(package.XlsxRejected, match="UNSUPPORTED_NUMBER_FORMAT"):
        loader._parse(percentage_workbook(0.476, "[Red]0.0%"), None, DEFAULT_LIMITS)


def test_builtin_static_print_and_filter_names_are_not_formulas():
    book = Workbook()
    sheet = book.active
    sheet.append(["Ay", "Faiz (%)"])
    sheet.append(["2024-01", 4])
    sheet.append(["2024-02", 5])
    sheet.auto_filter.ref = "A1:B3"
    sheet.print_area = "A1:B3"
    sheet.print_title_rows = "1:1"
    output = io.BytesIO()
    book.save(output)
    book.close()
    parsed = loader._parse(output.getvalue(), None, DEFAULT_LIMITS)
    assert parsed["rows"] == [["2024-01", 4], ["2024-02", 5]]


@pytest.mark.parametrize(
    "expression", ["OFFSET(Sheet!$A$1,0,0,2,2)", "'[external.xlsx]Sheet'!$A$1:$B$3", "1+1"]
)
def test_builtin_name_cannot_hide_formula_or_external_reference(expression):
    content = rewrite(
        SINGLE,
        {
            "xl/workbook.xml": lambda data: data.replace(
                b"<definedNames/>",
                (
                    '<definedNames><definedName name="_xlnm.Print_Area">'
                    + expression
                    + "</definedName></definedNames>"
                ).encode(),
            )
        },
    )
    with pytest.raises(package.XlsxRejected):
        loader._parse(content, None, DEFAULT_LIMITS)
